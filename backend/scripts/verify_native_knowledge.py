import asyncio, json, io, time
from pathlib import Path
from datetime import datetime, UTC
import httpx
from app.db.session import get_db_context
from app.db.models.user import User
from app.core.security import create_access_token, get_password_hash
from app.services.knowledge_index import parse_document, split_document, rank_chunks
from sqlalchemy import text


async def main():
    from app.core.config import settings

    assert settings.POSTGRES_DB.endswith("_review"), "Use a disposable database ending in _review"
    fixtures = []
    async with get_db_context() as db:
        for i in range(2):
            u = User(
                email=f"native-review-{i}@example.com",
                full_name=f"Native Review {i}",
                hashed_password=get_password_hash("NativeReview-2026!"),
                onboarding_completed_at=datetime.now(UTC),
            )
            db.add(u)
            await db.flush()
            fixtures.append(
                {
                    "id": str(u.id),
                    "email": u.email,
                    "password": "NativeReview-2026!",
                    "token": create_access_token(u.id),
                }
            )
    a, b = fixtures
    async with httpx.AsyncClient(base_url="http://localhost:38001", timeout=90) as c:

        def h(u):
            return {"Authorization": "Bearer " + u["token"]}

        async def req(method, path, u=a, **kw):
            return await c.request(method, "/api/v1/knowledge/" + path, headers=h(u), **kw)

        r = await req("POST", "bases", json={"name": "中文检索验收", "description": "独立测试资料"})
        assert r.status_code == 201, r.text
        kb = r.json()["id"]
        a["base_id"] = kb
        raw = "星河项目报销政策\n差旅餐饮每日上限为人民币 180 元，住宿每日上限为人民币 650 元。\n员工应在出差结束后 15 个工作日内提交报销申请。\n研发资料只允许项目成员访问。\n"
        r = await req(
            "POST",
            f"bases/{kb}/documents",
            files={"file": ("差旅管理办法.txt", raw.encode(), "text/plain")},
        )
        assert r.status_code == 201, r.text
        doc = r.json()["id"]
        a["document_id"] = doc
        dup = await req(
            "POST",
            f"bases/{kb}/documents",
            files={"file": ("另一名称.txt", raw.encode(), "text/plain")},
        )
        assert dup.json()["id"] == doc
        seen = set()
        ready = None
        for _ in range(100):
            r = await req("GET", f"bases/{kb}/documents")
            ready = r.json()[0]
            seen.add(ready["status"])
            if ready["status"] in ["ready", "failed"]:
                break
            await asyncio.sleep(0.5)
        assert ready["status"] == "ready", ready
        assert ready["chunk_count"] > 0
        print(
            "Upload → local parsing / embedding → ready; observed states:", sorted(seen), flush=True
        )
        r = await req(
            "POST",
            "search",
            json={"knowledge_base_ids": [kb], "query": "出差住宿一天最多能报多少钱？"},
        )
        assert r.status_code == 200, r.text
        hits = r.json()["items"]
        assert hits and "650" in hits[0]["content"], hits
        assert hits[0]["index"] == 1 and hits[0]["score_type"] == "rrf"
        print(
            "Chinese semantic paraphrase + hybrid retrieval + reference metadata: PASS", flush=True
        )
        assert (await req("GET", f"documents/{doc}/download")).content == raw.encode()
        assert len((await req("GET", f"documents/{doc}/chunks")).json()["items"]) > 0
        for method, path, kw in [
            ("GET", f"bases/{kb}/documents", {}),
            ("GET", f"documents/{doc}/download", {}),
            ("GET", f"documents/{doc}/chunks", {}),
            ("DELETE", f"documents/{doc}", {}),
            ("POST", f"documents/{doc}/retry", {}),
            ("DELETE", f"bases/{kb}", {}),
            ("POST", "search", {"json": {"knowledge_base_ids": [kb], "query": "住宿"}}),
            ("POST", f"bases/{kb}/documents", {"files": {"file": ("foreign.txt", b"foreign")}}),
        ]:
            r = await req(method, path, u=b, **kw)
            assert r.status_code == 404, (method, path, r.status_code, r.text)
        assert (await req("GET", "bases", u=b)).json()["items"] == []
        assert (await c.get("/api/v1/knowledge/bases")).status_code in [401, 403]
        print(
            "Two-account isolation: list, upload, search, retry, preview, download, deletion: PASS",
            flush=True,
        )
        # Word tables and PDF page metadata really pass through the background worker.
        from docx import Document
        import pymupdf

        word = Document()
        word.add_paragraph("月球实验室的会议室开放时间是上午九点至下午六点。")
        table = word.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "负责人"
        table.cell(0, 1).text = "林老师"
        stream = io.BytesIO()
        word.save(stream)
        pdf = pymupdf.open()
        pdf.new_page().insert_text((72, 72), "Lunar lab emergency phone is 8675309.")
        pdf_bytes = pdf.tobytes()
        pdf.close()
        for filename, data in [
            ("实验室.docx", stream.getvalue()),
            ("handbook.pdf", pdf_bytes),
            ("bad.pdf", b"not a pdf"),
        ]:
            r = await req("POST", f"bases/{kb}/documents", files={"file": (filename, data)})
            assert r.status_code == 201, r.text
        for _ in range(100):
            docs = (await req("GET", f"bases/{kb}/documents")).json()
            if all(d["status"] in ["ready", "failed"] for d in docs):
                break
            await asyncio.sleep(0.5)
        statuses = {d["filename"]: d["status"] for d in docs}
        assert (
            statuses["实验室.docx"] == "ready"
            and statuses["handbook.pdf"] == "ready"
            and statuses["bad.pdf"] == "failed"
        ), statuses
        failed = next(d for d in docs if d["filename"] == "bad.pdf")
        r = await req("POST", f"documents/{failed['id']}/retry")
        assert r.json()["status"] == "pending"
        r = await req("POST", "search", json={"knowledge_base_ids": [kb], "query": "8675309"})
        assert any(x["page"] == 1 and "8675309" in x["content"] for x in r.json()["items"])
        assert (
            await req("POST", "search", json={"knowledge_base_ids": [kb], "query": " "})
        ).status_code == 400
        assert (
            await req("POST", f"bases/{kb}/documents", files={"file": ("virus.exe", b"hello")})
        ).status_code == 400
        assert (await req("DELETE", f"documents/{failed['id']}")).status_code == 204
        print(
            "PDF pages, Word tables, invalid-file failure, retry, deletion and duplicate detection: PASS",
            flush=True,
        )
    Path("/tmp/native-fixtures.json").write_text(json.dumps(fixtures))
    Path("/tmp/native-fixtures.json").chmod(0o600)
    assert "你好" in parse_document("你好".encode("gb18030"), "中文.txt")[0][1]
    chunks = split_document([(1, "天地玄黄。" * 500)], size=120, overlap=20)
    assert all(len(c["content"]) <= 120 and c["page"] == 1 for c in chunks)
    print("Chinese encoding and chunk/page invariants: PASS", flush=True)


asyncio.run(main())
