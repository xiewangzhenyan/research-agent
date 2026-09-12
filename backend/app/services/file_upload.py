"""File upload service."""

import io
import logging
from typing import Any

from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import BadRequestError, NotFoundError
from app.db.models.chat_file import ChatFile
from app.repositories import chat_file as chat_file_repo
from app.services.file_storage import (
    ALLOWED_MIME_TYPES,
    MAX_UPLOAD_SIZE,
    classify_file,
    get_file_storage,
)

logger = logging.getLogger(__name__)


class FileUploadService:
    """Service for file upload validation, parsing, and persistence."""

    ALLOWED_MIME_TYPES = ALLOWED_MIME_TYPES
    MAX_UPLOAD_SIZE = MAX_UPLOAD_SIZE

    def __init__(self, db: AsyncSession):
        self.db = db

    @staticmethod
    def validate_upload(content_type: str | None, size: int) -> tuple[bool, str | None]:
        """Validate file type and size.

        Returns:
            Tuple of (is_valid, error_message).
        """
        if content_type not in ALLOWED_MIME_TYPES:
            return False, f"File type '{content_type}' is not supported."
        if size > MAX_UPLOAD_SIZE:
            return False, f"File too large. Maximum size is {MAX_UPLOAD_SIZE // (1024 * 1024)}MB."
        return True, None

    @staticmethod
    def classify_file(mime_type: str, filename: str) -> str:
        """Classify file type based on MIME type and extension."""
        return classify_file(mime_type, filename)

    async def parse_content(
        self,
        data: bytes,
        file_type: str,
        mime_type: str = "",
    ) -> str | None:
        """Parse file content based on file type.

        Returns extracted text content or None if parsing fails.
        """
        if file_type == "text":
            return self._parse_text_content(data, mime_type)
        elif file_type == "pdf":
            return self._parse_pdf_content(data)
        elif file_type == "docx":
            return self._parse_docx_content(data)
        return None

    @staticmethod
    def _parse_text_content(data: bytes, mime_type: str) -> str | None:
        """Extract text content from text-based files."""
        try:
            return data.decode("utf-8")
        except (UnicodeDecodeError, ValueError):
            return None

    @staticmethod
    def _parse_pdf_content(data: bytes) -> str | None:
        """Extract text from PDF using PyMuPDF."""
        try:
            import pymupdf

            doc: Any = pymupdf.open(stream=data, filetype="pdf")  # type: ignore[no-untyped-call,unused-ignore]
            texts = []
            for page in doc:
                blocks = page.get_text("blocks")
                for b in blocks:
                    if b[6] == 0:
                        text = b[4].strip()
                        if text:
                            texts.append(text)
                try:
                    tables = page.find_tables()
                    if tables and tables.tables:
                        for table in tables.tables:
                            df = table.to_pandas()
                            if not df.empty:
                                texts.append(df.to_markdown(index=False))
                except Exception:
                    pass
            doc.close()
            return "\n\n".join(texts) if texts else None
        except Exception as e:
            logger.warning("PDF parsing failed: %s", e)
            return None

    @staticmethod
    def _parse_docx_content(data: bytes) -> str | None:
        """Extract text from DOCX."""
        try:
            from docx import Document as DOCXDocument

            doc: Any = DOCXDocument(io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as e:
            logger.warning("DOCX parsing failed: %s", e)
            return None

    async def upload(
        self,
        *,
        user_id: Any,
        file_data: bytes,
        filename: str,
        content_type: str | None,
    ) -> ChatFile:
        """Validate, parse, persist, and record a chat file upload.

        Raises:
            BadRequestError: If file type or size is invalid.
        """
        is_valid, error = self.validate_upload(content_type, len(file_data))
        if not is_valid:
            raise BadRequestError(message=error or "Invalid file")

        file_type = self.classify_file(content_type or "", filename)
        parsed_content = await self.parse_content(file_data, file_type, content_type or "")
        if file_type in {"text", "pdf", "docx"} and not (parsed_content or "").strip():
            raise BadRequestError(
                message="无法提取文件文字，请上传包含可复制文字的 PDF、Word 或 UTF-8 文本。"
                "扫描件请先进行 OCR。 / No readable text found; use a text-based document."
            )

        storage = get_file_storage()
        storage_path = await storage.save(str(user_id), filename, file_data)

        return await self.create_chat_file(
            user_id=user_id,
            filename=filename,
            mime_type=content_type or "application/octet-stream",
            size=len(file_data),
            storage_path=storage_path,
            file_type=file_type,
            parsed_content=parsed_content,
        )

    def get_file_path(self, storage_path: str) -> str | None:
        """Resolve a storage path to an absolute filesystem path."""
        full_path = get_file_storage().get_full_path(storage_path)
        return str(full_path) if full_path is not None else None

    async def get_user_file(self, file_id: Any, user_id: Any) -> ChatFile:
        """Get a file by ID, verifying ownership.

        Raises:
            NotFoundError: If file does not exist or user has no access.
        """
        chat_file = await chat_file_repo.get_by_id(self.db, file_id)
        if not chat_file or str(chat_file.user_id) != str(user_id):
            raise NotFoundError(message="File not found")
        return chat_file

    async def create_chat_file(
        self,
        *,
        user_id: Any,
        filename: str,
        mime_type: str,
        size: int,
        storage_path: str,
        file_type: str,
        parsed_content: str | None = None,
    ) -> ChatFile:
        """Create a chat file record in the database."""
        return await chat_file_repo.create(
            self.db,
            user_id=user_id,
            filename=filename,
            mime_type=mime_type,
            size=size,
            storage_path=storage_path,
            file_type=file_type,
            parsed_content=parsed_content,
        )
