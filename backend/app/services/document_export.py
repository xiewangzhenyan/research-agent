"""Bounded, data-only Office rendering. No generated code, network or filesystem I/O."""

import asyncio
import hashlib
import re
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from typing import Literal
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from markdown_it import MarkdownIt
from pydantic import BaseModel, ConfigDict, Field

from app.core.exceptions import BadRequestError

DocumentFormat = Literal["md", "docx", "xlsx", "pptx"]
MIME_TYPES = {
    "md": "text/markdown",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
MAX_CONTENT = 100000
MAX_FILE = 2 * 1024**2
EPOCH = datetime(2000, 1, 1)
# Limit concurrent in-process template rendering (also used by the worker).
_render_slots = asyncio.Semaphore(2)


class ExportRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    format: DocumentFormat


class DocumentRequest(ExportRequest):
    title: str = Field(min_length=1, max_length=120)
    content: str = Field(min_length=1, max_length=MAX_CONTENT)


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    rows: list[list[str]] | None = None


def inline_text(tokens):
    parts = []
    links = []
    for token in tokens or []:
        if token.type == "link_open":
            links.append(token.attrGet("href") or "")
        elif token.type == "link_close":
            url = links.pop() if links else ""
            if url:
                parts.append(f" ({url})")
        elif token.type in {"text", "code_inline", "html_inline", "image"}:
            parts.append(token.content)
        elif token.type in {"softbreak", "hardbreak"}:
            parts.append("\n")
    return "".join(parts)


def parse_blocks(content):
    # Office currently keeps LaTeX as text. Protect delimiters, underscores and
    # backslashes from Markdown's emphasis/escape processing before restoring it.
    prefix = "LSPRAIMATH" + hashlib.sha256(content.encode()).hexdigest()[:16]
    while prefix in content:
        prefix += "X"
    math = {}

    pieces, cursor, missing_closers = [], 0, set()
    for match in re.finditer(r"\\\[|\\\(|\$\$|(?<![\\$])\$(?!\$)", content):
        if match.start() < cursor:
            continue
        opener = match.group()
        closer = {r"\[": r"\]", r"\(": r"\)"}.get(opener, opener)
        if closer in missing_closers:
            continue
        end = content.find(closer, match.end())
        if end == -1:
            # No later opener can find this missing closer either. Avoid repeated
            # scans of a large malformed document with thousands of unclosed maths.
            missing_closers.add(closer)
            continue
        if opener == "$" and "\n" in content[match.end() : end]:
            continue
        end += len(closer)
        key = f"{prefix}N{len(math)}END"
        math[key] = content[match.start() : end]
        pieces.extend([content[cursor : match.start()], key])
        cursor = end
    pieces.append(content[cursor:])
    protected = "".join(pieces)
    tokens = MarkdownIt("commonmark", {"html": False}).enable("table").parse(protected)

    def restore(text):
        return re.sub(prefix + r"N\d+END", lambda match: math[match.group()], text)

    if len(tokens) > 12000:
        raise BadRequestError(message="文档结构过多，请按章节拆分导出")
    blocks, rows, row = [], None, None
    list_depth, cell_count, table_count = 0, 0, 0
    for i, token in enumerate(tokens):
        if token.type == "table_open":
            rows = []
            table_count += 1
        elif token.type == "tr_open":
            row = []
        elif token.type == "tr_close":
            assert rows is not None and row is not None
            rows.append(row)
        elif token.type == "table_close":
            blocks.append(Block("table", rows=rows))
            rows = None
        elif token.type in {"bullet_list_open", "ordered_list_open"}:
            list_depth += 1
        elif token.type in {"bullet_list_close", "ordered_list_close"}:
            list_depth -= 1
        elif token.type == "inline":
            text = restore(inline_text(token.children))
            if rows is not None:
                assert row is not None
                cell_count += 1
                if len(text) > 4000:
                    raise BadRequestError(message="表格单元格过长，请拆分内容")
                row.append(text)
            else:
                previous = tokens[i - 1]
                kind = (
                    "heading"
                    if previous.type == "heading_open"
                    else "list"
                    if list_depth
                    else "text"
                )
                level = int(previous.tag[1:]) if kind == "heading" else list_depth
                blocks.append(Block(kind, text, level))
        elif token.type in {"fence", "code_block"}:
            blocks.append(Block("code", restore(token.content)))
    if cell_count > 10000 or table_count > 20 or len(blocks) > 2000:
        raise BadRequestError(message="内容超过文档结构限制，请按章节或表格拆分")
    return blocks


def stable_zip(content):
    """Normalize Office package timestamps so retries produce identical artifacts."""
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            data = source.read(name)
            if name == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]+",
                    rb"\g<1>2000-01-01T00:00:00Z",
                    data,
                )
            info = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, data)
    return output.getvalue()


def word_document(title, blocks):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    doc.core_properties.author = "LSPRAI"
    doc.core_properties.title = title
    doc.core_properties.created = doc.core_properties.modified = EPOCH
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    doc.add_heading(title, 0)
    for block in blocks:
        if block.kind == "table":
            rows = block.rows
            if len(rows[0]) > 16:
                raise BadRequestError(message="Word 表格最多 16 列，请拆分或改用 Excel")
            table = doc.add_table(rows=0, cols=len(rows[0]), style="Table Grid")
            for index, row in enumerate(rows):
                cells = table.add_row().cells
                for cell, text in zip(cells, row, strict=True):
                    cell.text = text
                    if index == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "E4F3EE")
                        cell._tc.get_or_add_tcPr().append(shade)
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
        elif block.kind == "heading":
            doc.add_heading(block.text, min(block.level, 6))
        elif block.kind == "list":
            doc.add_paragraph(block.text, style="List Bullet")
        else:
            paragraph = doc.add_paragraph(block.text)
            if block.kind == "code":
                for run in paragraph.runs:
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def spreadsheet(title, blocks):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.properties.creator = "LSPRAI"
    book.properties.title = title
    book.properties.created = book.properties.modified = EPOCH
    notes = book.active
    notes.title = "说明"
    notes.cell(1, 1, title).data_type = "s"

    def write(sheet, row, column, text, *, numeric=False):
        cell = sheet.cell(row, column)
        # All non-numeric content is explicitly text, including =,+,-,@ payloads.
        # Preserve leading-zero identifiers and numbers beyond Excel's precision.
        if (
            numeric
            and re.fullmatch(r"-?(?:0|[1-9]\d*)(?:\.\d+)?", text)
            and len(text.replace("-", "").replace(".", "")) <= 15
        ):
            cell.value = float(text) if "." in text else int(text)
        else:
            cell.value = text
            cell.data_type = "s"
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    table_index = 0
    for block in blocks:
        if block.kind == "table":
            table_index += 1
            sheet = book.create_sheet(f"表格 {table_index}")
            for row_index, row in enumerate(block.rows, 1):
                for col_index, text in enumerate(row, 1):
                    write(sheet, row_index, col_index, text, numeric=row_index > 1)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            write(notes, notes.max_row + 1, 1, f"[表格 {table_index}] 见同名工作表")
        else:
            # Excel silently truncates strings above 32767 chars; split first.
            for start in range(0, len(block.text), 4000):
                write(notes, notes.max_row + 1, 1, block.text[start : start + 4000])
    for sheet in book:
        for cell in sheet[1]:
            cell.font = Font(name="Microsoft YaHei", bold=True, color="143D34")
            cell.fill = PatternFill("solid", fgColor="E4F3EE")
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 65 if sheet is notes else 25
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def wrapped_lines(text, width=60):
    import unicodedata

    lines = []
    for original in text.split("\n"):
        line, used = "", 0
        for char in original:
            cost = 2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1
            if used + cost > width:
                lines.append(line)
                line, used = "", 0
            line += char
            used += cost
        lines.append(line)
    return lines


def presentation(title, blocks):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.core_properties.author = "LSPRAI"
    deck.core_properties.title = title
    deck.core_properties.created = deck.core_properties.modified = EPOCH

    def add_slide(heading, lines):
        if len(wrapped_lines(heading, 70)) > 2:
            raise BadRequestError(message="PPT 标题过长，请缩短标题或改为正文")
        if len(deck.slides) >= 60:
            raise BadRequestError(message="PPT 超过 60 页，请缩短内容或分批生成")
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for text, left, top, width, height, size, bold in (
            (heading, 0.9, 0.5, 11.5, 1.25, 28, True),
            ("\n".join(lines), 0.9, 1.9, 11.5, 4.8, 20, False),
            (f"LSPRAI · {len(deck.slides)}", 0.9, 7.0, 11.5, 0.3, 10, False),
        ):
            frame = slide.shapes.add_textbox(
                Inches(left), Inches(top), Inches(width), Inches(height)
            ).text_frame
            frame.word_wrap = True
            frame.text = text
            for paragraph in frame.paragraphs:
                paragraph.space_after = Pt(5)
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run.font.size = Pt(size)
                    run.font.bold = bold
                    run.font.color.rgb = RGBColor.from_string("174C40" if bold else "263442")

    # Text slides deliberately preserve table content as row text; Excel/Word retain grids.
    heading, pending = title, []

    def flush():
        nonlocal pending
        for start in range(0, len(pending), 10):
            add_slide(heading, pending[start : start + 10])
        pending = []

    for block in blocks:
        if block.kind == "heading":
            flush()
            heading = block.text
            if len(wrapped_lines(heading, 70)) > 2:
                raise BadRequestError(message="PPT 标题过长，请缩短标题或改为正文")
            pending = [""]
        else:
            text = (
                "\n".join(" | ".join(row) for row in block.rows)
                if block.kind == "table"
                else block.text
            )
            pending.extend(wrapped_lines(("• " if block.kind == "list" else "") + text))
    flush()
    if not deck.slides:
        add_slide(title, [])
    output = BytesIO()
    deck.save(output)
    return output.getvalue()


def render_document(request: DocumentRequest):
    # Reject XML-invalid controls instead of silently losing source characters.
    if re.search(
        r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", request.content + request.title
    ):
        raise BadRequestError(message="内容包含无法导出的控制字符，请清理后重试")
    if not request.content.strip() or not request.title.strip():
        raise BadRequestError(message="文档标题和正文不能为空")
    if request.format == "md":
        content = request.content.encode("utf-8")
    else:
        blocks = parse_blocks(request.content)
        renderer = {"docx": word_document, "xlsx": spreadsheet, "pptx": presentation}[
            request.format
        ]
        content = stable_zip(renderer(request.title, blocks))
    if len(content) > MAX_FILE:
        raise BadRequestError(message="文件超过 2 MiB，请拆分后生成")
    name = (
        re.sub(r"[^\w .-]", "_", request.title).replace("..", "_").strip(" ._")[:80] or "document"
    )
    return {
        "name": f"{name}.{request.format}",
        "mime_type": MIME_TYPES[request.format],
        "content": content,
        "size": len(content),
        "sha256": hashlib.sha256(content).hexdigest(),
    }


async def render_async(request):
    from anyio.to_thread import run_sync

    async with _render_slots:
        # Keep the event loop responsive; cancellation doesn't leave unbounded work behind.
        return await run_sync(render_document, request)
