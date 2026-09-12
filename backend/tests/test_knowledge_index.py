"""Regression tests for Chinese parsing, bounded chunks and hybrid ranking."""

import io
from itertools import pairwise

import pytest
from docx import Document

from app.services.knowledge_index import parse_document, rank_chunks, split_document


def test_chinese_text_encodings():
    for encoding in ("utf-8-sig", "gb18030"):
        assert parse_document("差旅报销规范".encode(encoding), "policy.txt") == [
            (None, "差旅报销规范")
        ]


def test_word_tables_are_searchable():
    doc = Document()
    doc.add_paragraph("项目政策")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "住宿上限"
    table.cell(0, 1).text = "650 元"
    buf = io.BytesIO()
    doc.save(buf)
    pages = parse_document(buf.getvalue(), "政策.docx")
    assert "住宿上限 | 650 元" in "\n".join(text for _, text in pages)


def test_blank_scanned_document_fails_instead_of_claiming_ready():
    import pymupdf

    doc = pymupdf.open()
    doc.new_page()
    with pytest.raises(ValueError, match="OCR"):
        parse_document(doc.tobytes(), "scan.pdf")
    doc.close()


def test_chunk_bounds_and_pages_preserve_all_text():
    text = "星河计划。" * 200
    chunks = split_document([(3, text)], size=120, overlap=20)
    assert all(len(c["content"]) <= 120 and c["page"] == 3 for c in chunks)
    assert [c["position"] for c in chunks] == list(range(len(chunks)))
    assert chunks[0]["content"].startswith(text[:30])
    assert chunks[-1]["content"].endswith(text[-30:])
    assert all(a["content"][-20:] == b["content"][:20] for a, b in pairwise(chunks))
    with pytest.raises(ValueError):
        split_document([(None, text)], size=10, overlap=10)


def test_hybrid_search_preserves_exact_identifier_and_document_diversity():
    chunks = [
        {
            "document_id": "a",
            "position": 0,
            "content": "报销政策：住宿每日 650 元",
            "embedding": [1, 0],
        },
        {"document_id": "a", "position": 1, "content": "报销政策附件", "embedding": [1, 0]},
        {"document_id": "b", "position": 0, "content": "专用标识 LSPR8675309", "embedding": [0, 1]},
        {"document_id": "c", "position": 0, "content": "无关内容", "embedding": [-1, 0]},
    ]
    results = rank_chunks("LSPR8675309", [1, 0], chunks)
    assert {c["document_id"] for c in results} == {"a", "b"}
    assert len(results) == 2
    assert all("embedding" not in c and c["score_type"] == "rrf" for c in results)
    assert rank_chunks("missing", [1, 0], chunks[-1:]) == []


def test_word_preserves_table_position_and_row_headers():
    from app.services.knowledge_index import parse_blocks, split_blocks

    doc = Document()
    doc.add_heading("报销标准", 1)
    doc.add_paragraph("以下标准适用于星河项目。")
    table = doc.add_table(rows=3, cols=2)
    for row, values in zip(
        table.rows, [("类别", "限额"), ("住宿", "650 元"), ("餐饮", "180 元")], strict=True
    ):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    doc.add_paragraph("表后备注：仅工作日适用。")
    buf = io.BytesIO()
    doc.save(buf)
    blocks = parse_blocks(buf.getvalue(), "制度.docx")
    text = "\n".join(b["text"] for b in blocks)
    assert text.index("星河项目") < text.index("650") < text.index("表后备注")
    chunks = split_blocks(blocks)
    row = next(c for c in chunks if "650" in c["content"])
    assert "类别 | 限额" in row["content"]
    assert row["location"] == {"section": "报销标准", "table": 1, "row": 2, "kind": "table_row"}
    assert row["page"] is None


def test_csv_quoted_multiline_rows_and_oversized_cells():
    from app.services.knowledge_index import parse_blocks, split_blocks

    raw = (
        '项目,说明,限额\r\n星河,"含税,含服务费\n工作日",650\r\n月球,"' + "详情" * 500 + '",180\r\n'
    )
    chunks = split_blocks(parse_blocks(raw.encode("utf-8-sig"), "预算.csv"))
    assert all(len(c["content"]) <= 450 and "项目 | 说明 | 限额" in c["content"] for c in chunks)
    row2 = [c for c in chunks if c["location"]["row"] == 2]
    assert len(row2) == 1 and "含税,含服务费\n工作日" in row2[0]["content"]
    assert all(c["page"] is None for c in chunks)


def test_adjacent_table_records_are_not_removed_as_overlap():
    chunks = [
        {
            "document_id": "a",
            "position": i,
            "location": {"row": i + 1},
            "content": f"限额 {i}",
            "embedding": [1, 0],
        }
        for i in range(3)
    ]
    assert len(rank_chunks("限额", [1, 0], chunks)) == 3


def test_table_header_is_context_not_a_separate_search_hit():
    from app.services.knowledge_index import parse_blocks, split_blocks

    chunks = split_blocks(parse_blocks("项目,补贴\n星云,每月 73 元".encode(), "补贴.csv"))
    assert len(chunks) == 1
    assert chunks[0]["location"]["row"] == 2
    assert "项目 | 补贴" in chunks[0]["content"] and "73" in chunks[0]["content"]


def test_adjacent_pdf_pages_remain_available_as_independent_evidence():
    chunks = [
        {
            "document_id": "paper",
            "position": i,
            "page": i + 1,
            "location": {},
            "content": text,
            "embedding": [1, 0],
        }
        for i, text in enumerate(
            ["灵敏度为 151 nm/RIU", "波长范围为 510 至 710 nm", "孵育时间为 18 分钟"]
        )
    ]
    hits = rank_chunks("灵敏度、波长范围和孵育时间", [1, 0], chunks)
    assert {h["page"] for h in hits} == {1, 2, 3}
    assert [h["index"] for h in hits] == [1, 2, 3]
    assert len(rank_chunks("灵敏度、波长范围和孵育时间", [1, 0], chunks, top_k=2)) == 2


def test_adjacent_word_sections_are_not_treated_as_overlap():
    from app.services.knowledge_index import parse_blocks, split_blocks

    doc = Document()
    doc.add_heading("实验条件", 1)
    doc.add_paragraph("合成实验的孵育时间为 18 分钟。")
    doc.add_heading("实验结果", 1)
    doc.add_paragraph("合成实验的灵敏度为 151 nm/RIU。")
    buf = io.BytesIO()
    doc.save(buf)
    chunks = [
        c | {"document_id": "word", "embedding": [1, 0]}
        for c in split_blocks(parse_blocks(buf.getvalue(), "study.docx"))
    ]
    hits = rank_chunks("合成实验条件和结果", [1, 0], chunks)
    assert {h["location"]["section"] for h in hits} == {"实验条件", "实验结果"}
    assert all(h["page"] is None for h in hits)


def test_same_page_overlap_rule_and_table_rows_keep_existing_behavior():
    chunks = [
        {
            "document_id": "paper",
            "position": i,
            "page": 4,
            "location": {},
            "content": "共同段落 " * 30,
            "embedding": [1, 0],
        }
        for i in range(3)
    ]
    hits = rank_chunks("共同段落", [1, 0], chunks)
    assert len(hits) == 2
    assert abs(hits[0]["position"] - hits[1]["position"]) > 1
