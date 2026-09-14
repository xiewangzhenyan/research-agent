"""Opt-in tests with a disposable DB; all model responses are local fakes."""

import asyncio
import json
import os
from io import BytesIO
from unittest.mock import patch
from uuid import UUID

import pytest
from docx import Document
from pydantic_ai.models.function import DeltaToolCall, FunctionModel

from app.api.routes.v1.chat_turns import download_answer
from app.core.config import settings
from app.core.exceptions import BadRequestError, NotFoundError
from app.db.models.agent_run import AgentRun
from app.db.models.conversation import Message
from app.db.models.user import User
from app.db.session import get_worker_db_context
from app.repositories.agent_run import AgentRunRepository
from app.schemas.project import ProjectWrite
from app.services.document_export import ExportRequest
from app.services.message_export import export_message
from app.services.project import ProjectService
from app.services.run_artifact import RunArtifactService
from app.worker.agent_runs import try_run
from tests.test_agent_runs_integration import get, users
from tests.test_durable_chat import state, submit

pytestmark = pytest.mark.skipif(os.getenv("RUN_TASK_DB_TESTS") != "1", reason="disposable DB only")


def test_parameters_persist_inherit_reset_and_isolate():
    async def check():
        async with users() as (owner, other):
            options = {"model": "gpt-4.1", "top_p": 0.45, "max_output_tokens": 2048}
            with patch.object(settings, "AI_MODEL_CONTROLS", {"gpt-4.1": ["temperature", "top_p"]}):
                first = await submit(owner, message="记住配置", generation=options)
                assert (await state(owner, first.conversation_id))["generation"] == options
                inherited = await submit(
                    owner, conversation_id=first.conversation_id, message="沿用参数"
                )
                assert inherited.effective_config["top_p"] == 0.45
                assert inherited.effective_config["temperature"] is None
                assert inherited.effective_config["max_output_tokens"] == 2048
                with pytest.raises(BadRequestError):
                    await submit(
                        owner,
                        conversation_id=first.conversation_id,
                        message="无效",
                        generation={"model": "gpt-4.1", "temperature": 0.2, "top_p": 0.5},
                    )
                assert (await state(owner, first.conversation_id))["generation"] == options
                reset = await submit(
                    owner, conversation_id=first.conversation_id, message="恢复默认", generation={}
                )
                assert (await state(owner, first.conversation_id))["generation"] == {}
                assert reset.effective_config["max_output_tokens"] == 8000
                assert (await get(owner, first.id)).effective_config["max_output_tokens"] == 2048
                new = await submit(owner, message="独立新会话")
                assert (await state(owner, new.conversation_id))["generation"] == {}
                with pytest.raises(NotFoundError):
                    await state(other, first.conversation_id)
                async with get_worker_db_context() as db:
                    project = await ProjectService(db, owner).save(ProjectWrite(name="隔离项目"))
                with pytest.raises(NotFoundError):
                    await state(owner, first.conversation_id, project.id)

    asyncio.run(check())


def test_export_checks_account_project_role_and_completion():
    async def check():
        async with users() as (owner, other):
            async with get_worker_db_context() as db:
                project = await ProjectService(db, owner).save(ProjectWrite(name="专用项目"))
            run = await submit(owner, project_id=project.id, message="导出报告")
            async with get_worker_db_context() as db:
                message = await db.get(Message, run.assistant_message_id)
                message.content = "# 结论\n\n有数据支持的结论。"
            for uid, pid in [(other, project.id), (owner, None)]:
                async with get_worker_db_context() as db:
                    with pytest.raises(NotFoundError):
                        await export_message(db, uid, pid, run.assistant_message_id, "docx")
            async with get_worker_db_context() as db:
                with pytest.raises(BadRequestError):
                    await export_message(db, owner, project.id, run.assistant_message_id, "docx")
                with pytest.raises(BadRequestError):
                    await export_message(db, owner, project.id, run.user_message_id, "docx")
                current = await db.get(AgentRun, run.id)
                current.status = "completed"
            async with get_worker_db_context() as db:
                user = await db.get(User, owner)
                response = await download_answer(
                    run.assistant_message_id, ExportRequest(format="docx"), user, db, project.id
                )
                assert "wordprocessingml" in response.headers["content-type"]
                assert response.headers["content-disposition"].startswith(
                    "attachment; filename*=UTF-8''"
                )
                assert "no-store" in response.headers["cache-control"]
                assert "有数据支持" in " ".join(
                    p.text for p in Document(BytesIO(response.body)).paragraphs
                )

    asyncio.run(check())


def test_document_tool_without_python_persists_reuses_and_survives_reopen():
    async def stream(messages, info):
        returned = [
            p
            for m in messages
            for p in m.parts
            if p.part_kind == "tool-return" and getattr(p, "tool_name", None) == "create_document"
        ]
        assert "run_python" not in [tool.name for tool in info.function_tools]
        if len(returned) >= 2:
            yield "报告已经生成，请使用文件卡片下载。"
        else:
            yield {
                0: DeltaToolCall(
                    name="create_document",
                    tool_call_id=f"document-{len(returned)}",
                    json_args=json.dumps(
                        {
                            "document": {
                                "format": "docx",
                                "title": "实验报告",
                                "content": "# 结果\n\n实验数据已经核对。",
                            }
                        },
                        ensure_ascii=False,
                    ),
                )
            }

    async def check():
        async with users() as (owner, other):
            run = await submit(owner, message="请生成一份 Word 实验报告")
            with patch(
                "app.agents.assistant._build_model",
                return_value=FunctionModel(stream_function=stream),
            ):
                await try_run(run.id)
            finished = await get(owner, run.id)
            assert finished.status == "completed", finished.error
            reopened = await state(owner, run.conversation_id)
            calls = reopened["messages"][1].tool_calls
            assert len(calls) == 2
            metadata = json.loads(calls[0].result)["artifacts"]
            assert metadata == json.loads(calls[1].result)["artifacts"]
            async with get_worker_db_context() as db:
                service = RunArtifactService(db, owner)
                assert len(await service.list(run.id)) == 1
                file = await service.get(run.id, UUID(metadata[0]["id"]))
                assert "实验数据" in " ".join(
                    p.text for p in Document(BytesIO(file.content)).paragraphs
                )
                with pytest.raises(NotFoundError):
                    await RunArtifactService(db, other).get(run.id, file.id)
                events = await AgentRunRepository(db).events(run.id, 0)
                assert any(e.kind == "document_created" for e in events)
                assert not any("content_base64" in json.dumps(e.data) for e in events)

    asyncio.run(check())
