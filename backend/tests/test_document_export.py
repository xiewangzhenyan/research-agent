"""Real Office packages, bounded work and spreadsheet injection regression."""

from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pydantic import ValidationError

from app.core.exceptions import BadRequestError
from app.services.document_export import DocumentRequest, parse_blocks, render_document

CONTENT = """# 实验记录

自由电子响应。引用资料 [1]，参数保留 00123。

| 标识 | 浓度 | 备注 |
| --- | --- | --- |
| 00123 | 1.25 | =HYPERLINK("https://invalid.test","x") |
| B | -2 | +cmd |

## 结论

- 需要复核。
- 不添加未知数据。

```python
print('data only')
```

[来源](https://example.invalid/paper)
"""


@pytest.mark.parametrize("format", ["md", "docx", "xlsx", "pptx"])
def test_real_unicode_document_is_deterministic_and_preserves_content(format):
    request = DocumentRequest(title="实验结果", content=CONTENT, format=format)
    file = render_document(request)
    assert file == render_document(request)
    assert file["name"] == f"实验结果.{format}"
    assert file["size"] == len(file["content"]) < 2 * 1024**2
    data = BytesIO(file["content"])
    if format == "md":
        assert data.getvalue().decode() == CONTENT
    elif format == "docx":
        doc = Document(data)
        assert doc.tables[0].cell(1, 0).text == "00123"
        assert "自由电子响应" in " ".join(p.text for p in doc.paragraphs)
        assert any("https://example.invalid/paper" in p.text for p in doc.paragraphs)
    elif format == "xlsx":
        book = load_workbook(data)
        sheet = book["表格 1"]
        assert sheet["A2"].value == "00123" and sheet["A2"].data_type == "s"
        assert sheet["B2"].value == 1.25 and sheet["B3"].value == -2
        assert sheet["C2"].value.startswith("=HYPERLINK") and sheet["C2"].data_type == "s"
        assert sheet["C3"].value == "+cmd" and sheet["C3"].data_type == "s"
        assert "结论" in str([cell.value for row in book["说明"] for cell in row])
    else:
        deck = Presentation(data)
        assert len(deck.slides) >= 2
        text = "\n".join(
            shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame
        )
        assert "00123" in text and "不添加未知数据" in text and "[1]" in text
    if format != "md":
        with ZipFile(BytesIO(file["content"])) as archive:
            assert not any("vbaProject" in name for name in archive.namelist())


def test_spreadsheet_title_is_text_and_long_paragraph_is_not_silently_truncated():
    request = DocumentRequest(title="=1+1", content="a" * 40000, format="xlsx")
    book = load_workbook(BytesIO(render_document(request)["content"]))
    assert book.active["A1"].data_type == "s"
    assert sum(len(row[0].value) for row in list(book.active)[1:]) == 40000


def test_no_network_images_or_generated_code_execution():
    content = "![外部图片](https://example.invalid/private)\n\n<script>fetch('/secret')</script>"
    file = render_document(DocumentRequest(title="../unsafe/path", content=content, format="docx"))
    assert "/" not in file["name"] and ".." not in file["name"]
    with ZipFile(BytesIO(file["content"])) as archive:
        assert not any(name.startswith("word/media/") for name in archive.namelist())
    assert "外部图片" in str(parse_blocks(content))


@pytest.mark.parametrize("content", ["x" * 100001, ""])
def test_rejects_unbounded_or_empty_content(content):
    with pytest.raises(ValidationError):
        DocumentRequest(title="test", content=content, format="docx")


def test_rejects_control_characters_large_cells_and_too_many_slides():
    for content, format in [
        ("bad\x00text", "docx"),
        ("| a |\n| - |\n|" + "x" * 4001 + "|", "xlsx"),
        ("长文" * 25000, "pptx"),
    ]:
        with pytest.raises(BadRequestError):
            render_document(DocumentRequest(title="test", content=content, format=format))


def test_office_preserves_latex_without_markdown_eating_delimiters_or_operators():
    formula = (
        r"\[\varepsilon(\omega)=\varepsilon_\infty-\frac{\omega_p^2}{\omega^2+i\gamma\omega}\]"
    )
    inline = r"$x_a + x_b$"
    request = DocumentRequest(title="Drude 模型", content=formula + "\n\n" + inline, format="docx")
    doc = Document(BytesIO(render_document(request)["content"]))
    paragraphs = [p.text for p in doc.paragraphs]
    assert formula in paragraphs and inline in paragraphs
