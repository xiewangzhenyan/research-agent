"""Preview resource boundaries, receipt isolation and actual parser equivalence."""

import asyncio
import fcntl
import io
from datetime import UTC, datetime, timedelta
from unittest.mock import AsyncMock
from uuid import uuid4

import jwt
import pytest

from app.core.config import settings
from app.core.exceptions import BadRequestError
from app.schemas.knowledge import ChunkingConfig
from app.services import knowledge_preview as preview
from app.services.knowledge_index import parse_with_report, split_blocks


def fixtures():
    from docx import Document

    word = Document()
    word.add_heading("操作指引", level=1)
    word.add_paragraph("先检查电源,再检查连接。" * 70)
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "限额"
    table.cell(1, 0).text = "住宿"
    table.cell(1, 1).text = "620元"
    buffer = io.BytesIO()
    word.save(buffer)
    import fitz

    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "General handbook. Travel limit is 620.")
    return [
        ("guide.txt", ("通用指南:先检查电源。" * 80).encode()),
        ("guide.md", ("# 操作指南\n\n保存后再关闭。" * 60).encode()),
        ("table.csv", '项目,限额\n住宿,620元\n"带换行\n内容",100元'.encode()),
        ("guide.docx", buffer.getvalue()),
        ("guide.pdf", pdf.tobytes()),
    ]


@pytest.mark.anyio
@pytest.mark.parametrize("filename,data", fixtures(), ids=[f[0] for f in fixtures()])
async def test_preview_reuses_actual_parser_and_splitter(filename, data):
    config = ChunkingConfig(chunk_size=300, chunk_overlap=0)
    result = await preview.run_preview(data, filename, config)
    blocks, report = parse_with_report(data, filename)
    chunks = split_blocks(blocks, size=300, overlap=0)
    assert result["items"] == chunks
    assert result["chunk_count"] == len(chunks)
    assert result["parse_report"] == report
    assert not result["truncated"]
    assert result["max_length"] == max(len(c["content"]) for c in chunks)
    assert all("embedding" not in item for item in result["items"])


@pytest.mark.anyio
async def test_truncation_only_limits_display_and_errors_are_actionable():
    result = await preview.run_preview(
        b"a" * 30_000, "long.txt", ChunkingConfig(chunk_size=300, chunk_overlap=0)
    )
    assert result["chunk_count"] == 100 and len(result["items"]) == 80
    assert result["truncated"] and result["mean_length"] == 300
    with pytest.raises(BadRequestError, match="无法读取文字"):
        await preview.run_preview(b"   ", "empty.txt", ChunkingConfig())


@pytest.mark.anyio
async def test_busy_lock_rejects_without_starting_another_process(monkeypatch):
    spawn = AsyncMock()
    monkeypatch.setattr(asyncio, "create_subprocess_exec", spawn)
    with open("/tmp/agent-knowledge-preview.lock", "a") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
        with pytest.raises(BadRequestError, match="服务忙"):
            await preview.run_preview(b"text", "a.txt", ChunkingConfig())
    spawn.assert_not_called()


@pytest.mark.anyio
@pytest.mark.parametrize("cancel", [False, True])
async def test_timeout_and_cancellation_reap_child_and_release_lock(monkeypatch, cancel):
    class Child:
        returncode = None
        killed = False
        waited = False

        async def communicate(self, data):
            if cancel:
                raise asyncio.CancelledError
            await asyncio.sleep(10)

        def kill(self):
            self.killed = True

        async def wait(self):
            self.waited = True
            self.returncode = -9

    child = Child()
    monkeypatch.setattr(asyncio, "create_subprocess_exec", AsyncMock(return_value=child))
    monkeypatch.setattr(preview, "PREVIEW_TIMEOUT", 0.01)
    with pytest.raises(asyncio.CancelledError if cancel else BadRequestError):
        await preview.run_preview(b"text", "a.txt", ChunkingConfig())
    assert child.killed and child.waited
    with open("/tmp/agent-knowledge-preview.lock", "a") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)


def test_confirmation_receipt_bound_to_owner_base_file_and_version():
    owner, base = uuid4(), uuid4()
    config = ChunkingConfig(chunk_size=300, chunk_overlap=0)
    token = preview.issue_receipt(owner, base, "a.txt", b"text", config)
    assert preview.verify_receipt(token, owner, base, "a.txt", b"text") == config
    for args in [
        (uuid4(), base, "a.txt", b"text"),
        (owner, uuid4(), "a.txt", b"text"),
        (owner, base, "b.txt", b"text"),
        (owner, base, "a.txt", b"changed"),
    ]:
        with pytest.raises(BadRequestError):
            preview.verify_receipt(token, *args)
    claims = jwt.decode(token, settings.SECRET_KEY, algorithms=[settings.ALGORITHM])
    for change in [
        {"exp": datetime.now(UTC) - timedelta(seconds=1)},
        {"type": "access"},
        {"version": "old-parser"},
        {"config": {"chunk_size": 900}},
    ]:
        invalid = jwt.encode(
            {**claims, **change}, settings.SECRET_KEY, algorithm=settings.ALGORITHM
        )
        with pytest.raises(BadRequestError):
            preview.verify_receipt(invalid, owner, base, "a.txt", b"text")
    with pytest.raises(BadRequestError):
        preview.verify_receipt(token + "invalid", owner, base, "a.txt", b"text")


@pytest.mark.parametrize(
    "name,data", [("a.exe", b"text"), ("a.txt", b""), ("a.txt", b"a" * (10 * 1024 * 1024 + 1))]
)
def test_preview_file_limits(name, data):
    with pytest.raises(BadRequestError):
        preview.validate_file(name, data)
